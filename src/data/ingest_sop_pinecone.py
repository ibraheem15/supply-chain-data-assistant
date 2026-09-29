"""Ingest SOP documents from data/policy/ into a Pinecone vector index.

Incremental: a content hash per source file (cached locally) decides whether a file
is new, changed, unchanged, or deleted since the last run, so reruns only touch what
actually changed instead of re-embedding everything every time.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import logging
import sys
import time
from collections.abc import Callable, Iterator
from pathlib import Path
from typing import Any

from langchain_core.documents import Document
from langchain_core.embeddings import Embeddings
from langchain_text_splitters import MarkdownHeaderTextSplitter, RecursiveCharacterTextSplitter

ROOT = Path(__file__).resolve().parents[2]
POLICY_DIR = ROOT / "data" / "policy"
CACHE_FILE = ROOT / ".cache" / "sop_ingest_cache.json"
VECTORDB_ENV_FILE = ROOT / "envs" / "vectorDB.env"
LLM_ENV_FILE = ROOT / "envs" / "LLM.env"

NO_SPLIT_SUFFIXES = {".csv", ".xlsx"}
CHUNK_SIZE = 400
CHUNK_OVERLAP = 60
BATCH_SIZE = 100

MARKDOWN_HEADERS = [("#", "Header 1"), ("##", "Header 2"), ("###", "Header 3")]

LOCAL_EMBEDDING_MODEL = "sentence-transformers/all-MiniLM-L6-v2"
LOCAL_EMBEDDING_DIMENSION = 384
GEMINI_EMBEDDING_MODEL = "models/text-embedding-004"
GEMINI_EMBEDDING_DIMENSION = 768

LOGGER = logging.getLogger("sop_ingest")


class IngestError(Exception):
    """An expected, actionable ingestion failure."""


# ==========================================
# 1. PATH RESOLUTION & SETUP
# ==========================================


def read_env_file(path: Path) -> dict[str, str]:
    """Read simple KEY=VALUE entries without logging their values."""
    values: dict[str, str] = {}
    if not path.exists():
        return values
    for line in path.read_text(encoding="utf-8").splitlines():
        stripped = line.strip()
        if not stripped or stripped.startswith("#") or "=" not in stripped:
            continue
        key, value = stripped.split("=", maxsplit=1)
        values[key.strip()] = value.strip().strip("'\"")
    return values


def load_settings() -> dict[str, str]:
    """Merge vectorDB.env and LLM.env; the process environment wins over both."""
    merged = {**read_env_file(VECTORDB_ENV_FILE), **read_env_file(LLM_ENV_FILE)}
    import os

    merged.update({k: v for k, v in os.environ.items() if k in merged or k.startswith(("PINECONE_", "GEMINI_", "EMBEDDING_"))})
    return merged


def _require(settings: dict[str, str], key: str) -> str:
    value = settings.get(key)
    if not value:
        raise IngestError(f"{key} is missing. Set it in vectorDB.env, LLM.env, or the process environment.")
    return value


# ==========================================
# 2. PINECONE PROVISIONING
# ==========================================


def ensure_pinecone_index(pc: Any, name: str, dimension: int, cloud: str, region: str) -> None:
    from pinecone import ServerlessSpec

    existing = {index["name"]: index for index in pc.list_indexes()}
    current = existing.get(name)

    if current is not None and current["dimension"] != dimension:
        LOGGER.warning(
            "Index %r has dimension %s but the configured embedding model produces %s; recreating.",
            name, current["dimension"], dimension,
        )
        pc.delete_index(name)
        current = None

    if current is None:
        pc.create_index(
            name=name,
            dimension=dimension,
            metric="cosine",
            spec=ServerlessSpec(cloud=cloud, region=region),
        )

    deadline = time.monotonic() + 60
    while time.monotonic() < deadline:
        if pc.describe_index(name).status["ready"]:
            return
        time.sleep(1)
    raise IngestError(f"Pinecone index {name!r} did not become ready within 60s.")


def build_embeddings(provider: str, settings: dict[str, str]) -> tuple[Embeddings, int]:
    if provider == "local":
        from langchain_huggingface import HuggingFaceEmbeddings

        return HuggingFaceEmbeddings(model_name=LOCAL_EMBEDDING_MODEL), LOCAL_EMBEDDING_DIMENSION
    if provider == "gemini":
        from langchain_google_genai import GoogleGenerativeAIEmbeddings

        api_key = _require(settings, "GEMINI_API_KEY")
        return (
            GoogleGenerativeAIEmbeddings(model=GEMINI_EMBEDDING_MODEL, google_api_key=api_key),
            GEMINI_EMBEDDING_DIMENSION,
        )
    raise IngestError(f"Unknown EMBEDDING_PROVIDER {provider!r}; expected 'local' or 'gemini'.")


# ==========================================
# 4. ROBUST POLYMORPHIC PARSER
# ==========================================


def _base_metadata(path: Path) -> dict[str, Any]:
    return {"source": str(path.relative_to(ROOT)), "file_type": path.suffix.lower().lstrip(".")}


def parse_markdown(path: Path) -> list[Document]:
    text = path.read_text(encoding="utf-8")
    splitter = MarkdownHeaderTextSplitter(headers_to_split_on=MARKDOWN_HEADERS, strip_headers=False)
    base = _base_metadata(path)
    sections = splitter.split_text(text)
    for section in sections:
        section.metadata = {**base, **section.metadata}
    return sections


def parse_text(path: Path) -> list[Document]:
    return [Document(page_content=path.read_text(encoding="utf-8"), metadata=_base_metadata(path))]


def parse_pdf(path: Path) -> list[Document]:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    documents = []
    for page_number, page in enumerate(reader.pages, start=1):
        text = page.extract_text() or ""
        if text.strip():
            documents.append(Document(page_content=text, metadata=_base_metadata(path) | {"page": page_number}))
    if not documents:
        raise IngestError(f"No extractable text in PDF: {path}")
    return documents


def parse_csv(path: Path) -> list[Document]:
    documents = []
    with path.open("r", encoding="utf-8-sig", newline="") as source:
        for row_number, row in enumerate(csv.DictReader(source), start=1):
            content = "\n".join(f"{key}: {value}" for key, value in row.items())
            documents.append(Document(page_content=content, metadata=_base_metadata(path) | {"row": row_number}))
    if not documents:
        raise IngestError(f"CSV has no data rows: {path}")
    return documents


def parse_xlsx(path: Path) -> list[Document]:
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    documents = []
    for sheet in workbook.worksheets:
        rows = sheet.iter_rows(values_only=True)
        header = next(rows, None)
        if header is None:
            continue
        for row_number, row in enumerate(rows, start=1):
            content = "\n".join(f"{col}: {value}" for col, value in zip(header, row) if value is not None)
            if content:
                documents.append(
                    Document(page_content=content, metadata=_base_metadata(path) | {"sheet": sheet.title, "row": row_number})
                )
    if not documents:
        raise IngestError(f"Workbook has no data rows: {path}")
    return documents


PARSERS: dict[str, Callable[[Path], list[Document]]] = {
    ".md": parse_markdown,
    ".txt": parse_text,
    ".pdf": parse_pdf,
    ".csv": parse_csv,
    ".xlsx": parse_xlsx,
}


def parse_file(path: Path) -> list[Document]:
    parser = PARSERS.get(path.suffix.lower())
    if parser is None:
        raise IngestError(f"Unsupported file type: {path.suffix}")
    return parser(path)


def split_documents(documents: list[Document]) -> list[Document]:
    splitter = RecursiveCharacterTextSplitter(chunk_size=CHUNK_SIZE, chunk_overlap=CHUNK_OVERLAP)
    return splitter.split_documents(documents)


# ==========================================
# 5. INCREMENTAL PIPELINE WITH BATCHING
# ==========================================


def _batched(items: list[Any], size: int) -> Iterator[list[Any]]:
    for offset in range(0, len(items), size):
        yield items[offset : offset + size]


def discover_files() -> list[Path]:
    return sorted(p for p in POLICY_DIR.rglob("*") if p.suffix.lower() in PARSERS)


def plan_sync(files: list[Path], cache: dict[str, dict[str, Any]]) -> dict[str, Any]:
    """Diff files on disk against the cache; returns pending chunks/ids/deletes and counts."""
    seen_paths: set[str] = set()
    pending_docs: list[Document] = []
    pending_ids: list[str] = []
    pending_deletes: list[str] = []
    counts = {"new": 0, "changed": 0, "unchanged": 0, "deleted": 0}

    for path in files:
        rel = str(path.relative_to(ROOT))
        seen_paths.add(rel)
        digest = hashlib.sha256(path.read_bytes()).hexdigest()
        cached = cache.get(rel)

        if cached and cached["hash"] == digest:
            counts["unchanged"] += 1
            continue

        try:
            documents = parse_file(path)
        except IngestError as exc:
            LOGGER.error("Skipping %s: %s", rel, exc)
            continue

        chunks = documents if path.suffix.lower() in NO_SPLIT_SUFFIXES else split_documents(documents)
        ids = [f"{rel}::{i}" for i in range(len(chunks))]

        if cached:
            pending_deletes.extend(cached["chunk_ids"])
            counts["changed"] += 1
            LOGGER.info("Changed %s -> %d chunk(s)", rel, len(chunks))
        else:
            counts["new"] += 1
            LOGGER.info("New %s -> %d chunk(s)", rel, len(chunks))

        pending_docs.extend(chunks)
        pending_ids.extend(ids)
        cache[rel] = {"hash": digest, "chunk_ids": ids}

    for rel in [rel for rel in cache if rel not in seen_paths]:
        pending_deletes.extend(cache[rel]["chunk_ids"])
        del cache[rel]
        counts["deleted"] += 1
        LOGGER.info("Deleted %s", rel)

    return {"docs": pending_docs, "ids": pending_ids, "deletes": pending_deletes, "counts": counts}


# ==========================================
# 6. SYNC HASH CACHE
# ==========================================


def load_cache(path: Path) -> dict[str, dict[str, Any]]:
    if not path.exists():
        return {}
    return json.loads(path.read_text(encoding="utf-8"))


def save_cache(path: Path, cache: dict[str, dict[str, Any]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    tmp_path = path.with_suffix(path.suffix + ".tmp")
    tmp_path.write_text(json.dumps(cache, indent=2), encoding="utf-8")
    tmp_path.replace(path)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Sync data/policy/ SOP documents into a Pinecone index.")
    parser.add_argument(
        "--dry-run",
        action="store_true",
        help="Report what would change without contacting Pinecone or the embedding model.",
    )
    return parser.parse_args()


def main() -> int:
    logging.basicConfig(level=logging.INFO, format="%(levelname)s: %(message)s")
    args = parse_args()
    try:
        files = discover_files()
        if not files:
            LOGGER.warning("No supported files found under %s.", POLICY_DIR)

        cache = load_cache(CACHE_FILE)
        plan = plan_sync(files, cache)
        counts = plan["counts"]
        LOGGER.info(
            "Scan complete: %d new, %d changed, %d unchanged, %d deleted.",
            counts["new"], counts["changed"], counts["unchanged"], counts["deleted"],
        )

        if args.dry_run:
            LOGGER.info("Dry run complete; Pinecone was not contacted.")
            return 0

        if not plan["docs"] and not plan["deletes"]:
            LOGGER.info("Nothing to sync.")
            return 0

        settings = load_settings()
        provider = settings.get("EMBEDDING_PROVIDER", "local")
        embeddings, dimension = build_embeddings(provider, settings)

        from pinecone import Pinecone
        from langchain_pinecone import PineconeVectorStore

        pc = Pinecone(api_key=_require(settings, "PINECONE_API_KEY"))
        index_name = settings.get("PINECONE_INDEX_NAME", "supply-chain-sops")
        ensure_pinecone_index(
            pc, index_name, dimension,
            cloud=settings.get("PINECONE_CLOUD", "aws"),
            region=settings.get("PINECONE_REGION", "us-east-1"),
        )
        vector_store = PineconeVectorStore(index=pc.Index(index_name), embedding=embeddings)

        for batch in _batched(plan["deletes"], BATCH_SIZE):
            vector_store.delete(ids=batch)
        for doc_batch, id_batch in zip(_batched(plan["docs"], BATCH_SIZE), _batched(plan["ids"], BATCH_SIZE)):
            vector_store.add_documents(doc_batch, ids=id_batch)

        save_cache(CACHE_FILE, cache)
        LOGGER.info("Synced %d chunk(s) to Pinecone index %r.", len(plan["docs"]), index_name)
        return 0
    except IngestError as exc:
        LOGGER.error("%s", exc)
        return 1


if __name__ == "__main__":
    sys.exit(main())
